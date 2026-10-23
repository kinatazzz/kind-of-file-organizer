from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile


TEXT_EXTENSIONS = {".txt", ".csv", ".md", ".json", ".log"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".xlsx"}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_TEXT_CHARS = 12000
MAX_PDF_PAGES = 4
MAX_SHEETS = 3
MAX_ROWS = 30
MAX_COLUMNS = 20


def check_archive(file_path: Path) -> None:
    with ZipFile(file_path) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_ARCHIVE_BYTES:
            raise ValueError("Слишком большой распакованный документ.")


def read_pdf(file_path: Path) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
    except ImportError as error:
        raise ValueError("Для PDF установите зависимости из requirements.txt.") from error

    try:
        with file_path.open("rb") as file:
            reader = PdfReader(file)
            if reader.is_encrypted:
                raise ValueError("PDF защищён паролем.")
            pages = reader.pages[:MAX_PDF_PAGES]
            return "\n".join(page.extract_text() or "" for page in pages)
    except PdfReadError as error:
        raise ValueError("Не удалось прочитать PDF.") from error


def read_docx(file_path: Path) -> str:
    check_archive(file_path)
    with ZipFile(file_path) as archive:
        try:
            document = archive.read("word/document.xml")
        except KeyError as error:
            raise ValueError("В DOCX отсутствует основной документ.") from error

    root = ElementTree.fromstring(document)
    return " ".join(
        element.text or ""
        for element in root.iter()
        if element.tag.rsplit("}", 1)[-1] == "t"
    )


def read_xlsx(file_path: Path) -> str:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as error:
        raise ValueError("Для XLSX установите зависимости из requirements.txt.") from error

    check_archive(file_path)
    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
    except (InvalidFileException, KeyError) as error:
        raise ValueError("Не удалось прочитать XLSX.") from error

    parts = []
    try:
        # берём начало первых листов
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            parts.append(sheet.title)
            for row in sheet.iter_rows(
                max_row=MAX_ROWS, max_col=MAX_COLUMNS, values_only=True
            ):
                parts.append(" | ".join(str(value) for value in row if value is not None))
                if sum(len(part) for part in parts) >= MAX_TEXT_CHARS:
                    return "\n".join(parts)
    finally:
        workbook.close()

    return "\n".join(parts)


def read_content(file_path: Path) -> str:
    if file_path.is_symlink():
        raise ValueError("Ссылки не отправляются на анализ.")
    if file_path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Файл превышает лимит анализа 10 МиБ.")

    extension = file_path.suffix.casefold()
    try:
        if extension in TEXT_EXTENSIONS:
            data = file_path.read_bytes()
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = data.decode("cp1251")
        elif extension == ".pdf":
            text = read_pdf(file_path)
        elif extension == ".docx":
            text = read_docx(file_path)
        elif extension == ".xlsx":
            text = read_xlsx(file_path)
        else:
            raise ValueError("Этот формат пока не поддерживается для анализа.")
    except (BadZipFile, ElementTree.ParseError, RuntimeError) as error:
        raise ValueError("Документ повреждён или защищён паролем.") from error

    # ограничиваем передаваемый фрагмент
    text = text.strip()[:MAX_TEXT_CHARS]
    if not text:
        raise ValueError("В документе нет доступного текста.")
    return text
